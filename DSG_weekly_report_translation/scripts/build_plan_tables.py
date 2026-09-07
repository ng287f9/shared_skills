#!/usr/bin/env python
"""CW36 test build: replace the Short Term Flight Test Plan screenshot tables
(slides 7/8) with native bilingual tables OCR'd from the EMF images; split
SN1004's CW36+CW37 combined table onto two slides (8 = CW36, new slide 9 = CW37).

Conventions (per user + ft-plan-translation): Arial 8pt, EN bold + CN regular in
the SAME color as the source line (red/blue/black/white-on-bar), table width =
banner image width, narrow label column, remaining 5 columns equal.
[bracket] crew-code lines and LFTE:xxx colon lines stay English-only; the
25/100/200FH inspection lines get ONE merged Chinese line.

Usage: python build_plan_tables_cw36.py <in.pptx> <out.pptx>
"""
import sys, os, zipfile, shutil, re
from pptx import Presentation
from pptx.util import Emu
from lxml import etree
sys.stdout.reconfigure(encoding='utf-8')

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'

GREY='C0C0C0'; PURPLE='D071FF'; ORANGE='FFC000'; SELGREY='BFBFBF'
GREEN='00B050'; RED='FF0000'; BLUE='0000FF'; BLACK='000000'; WHITE='FFFFFF'

def tag(e): return e.tag.split('}')[-1] if '}' in e.tag else e.tag

def make_para(text, color, bold, lang):
    p = etree.Element(A+'p')
    pPr = etree.SubElement(p, A+'pPr')
    pPr.set('marL','0'); pPr.set('indent','0'); pPr.set('algn','ctr')
    ln = etree.SubElement(pPr, A+'lnSpc'); etree.SubElement(ln, A+'spcPct').set('val','95000')
    for tn in ('spcBef','spcAft'):
        e = etree.SubElement(pPr, A+tn); etree.SubElement(e, A+'spcPts').set('val','0')
    etree.SubElement(pPr, A+'buNone')
    r = etree.SubElement(p, A+'r')
    rPr = etree.SubElement(r, A+'rPr')
    rPr.set('lang', 'zh-CN' if lang=='cn' else 'en-US'); rPr.set('altLang','en-US')
    rPr.set('sz','800'); rPr.set('dirty','0')
    if bold: rPr.set('b','1')
    sf = etree.SubElement(rPr, A+'solidFill'); etree.SubElement(sf, A+'srgbClr').set('val', color)
    la = etree.SubElement(rPr, A+'latin'); la.set('typeface','Arial'); la.set('pitchFamily','34'); la.set('charset','0')
    ea = etree.SubElement(rPr, A+'ea'); ea.set('typeface','Arial'); ea.set('pitchFamily','34'); ea.set('charset','0')
    etree.SubElement(r, A+'t').text = text
    return p

def set_cell(cell, spec, nrows):
    """spec: dict(fill=..., paras=[(text,color,bold,lang), ...])"""
    tc = cell._tc
    txBody = tc.find(A+'txBody')
    for p in txBody.findall(A+'p'): txBody.remove(p)
    for (text, color, bold, lang) in spec['paras']:
        txBody.append(make_para(text, color, bold, lang))
    if not spec['paras']:
        txBody.append(make_para('', BLACK, False, 'en'))
    tcPr = tc.find(A+'tcPr')
    if tcPr is None: tcPr = etree.SubElement(tc, A+'tcPr')
    tcPr.set('marL','27432'); tcPr.set('marR','27432'); tcPr.set('marT','9000'); tcPr.set('marB','9000')
    tcPr.set('anchor','ctr')
    for old in tcPr.findall(A+'solidFill'): tcPr.remove(old)
    if spec.get('fill'):
        sf = etree.SubElement(tcPr, A+'solidFill'); etree.SubElement(sf, A+'srgbClr').set('val', spec['fill'])
    # thin black borders (outer 12700 / inner 6350)
    ri, ci = spec['pos']
    for side, outer in (('lnL', ci==0), ('lnR', ci==5), ('lnT', ri==0), ('lnB', ri==nrows-1)):
        for old in tcPr.findall(A+side): tcPr.remove(old)
        ln = etree.SubElement(tcPr, A+side)
        ln.set('w', '12700' if outer else '6350')
        ln.set('cap','flat'); ln.set('cmpd','sng'); ln.set('algn','ctr')
        sf = etree.SubElement(ln, A+'solidFill'); etree.SubElement(sf, A+'srgbClr').set('val','000000')
        etree.SubElement(ln, A+'prstDash').set('val','solid')
        etree.SubElement(ln, A+'round')

def build_table(slide, rows, x, y, w, label_w):
    """rows: list of dicts(h=emu_height, cells=[spec]*6). Builds 6-col table."""
    nrows = len(rows)
    col_w = [int(label_w)] + [int((w-label_w)/5)]*4 + [int(w-label_w-int((w-label_w)/5)*4)]
    gf = slide.shapes.add_table(nrows, 6, Emu(x), Emu(y), Emu(w), sum(r['h'] for r in rows))
    tbl = gf.table
    for ci, cw in enumerate(col_w): tbl.columns[ci].width = Emu(cw)
    trs = tbl._tbl.findall(A+'tr')
    for ri, r in enumerate(rows): trs[ri].set('h', str(r['h']))
    for ri, r in enumerate(rows):
        for ci in range(6):
            spec = r['cells'][ci]
            spec['pos'] = (ri, ci)
            set_cell(tbl.cell(ri, ci), spec, nrows)
    tblPr = tbl._tbl.find(A+'tblPr')
    for c in list(tblPr): tblPr.remove(c)
    for k in list(tblPr.attrib): del tblPr.attrib[k]
    return tbl

def L(text, fill=GREY):  # label cell
    return dict(fill=fill, paras=[(text, BLACK, True, 'en')])
def A_hdr(text, fill):   # activity header
    return dict(fill=fill, paras=[(text, BLACK, True, 'en'), (CN_ACT[text], BLACK, False, 'cn')])
def det(pairs):          # detail cell: pairs = [(en, cn, color)]; cn None/str/list
    paras = []
    for en, cn, color in pairs:
        for line in en.split('\n'):
            paras.append((line, color, True, 'en'))
        if cn is None: continue
        if isinstance(cn, list):
            for c in cn: paras.append((c, color, False, 'cn'))
        else:
            paras.append((cn, color, False, 'cn'))
    return dict(fill=WHITE, paras=paras)
def bar(text, cn, fill):  # executed/cancelled status bar
    paras = []
    if text: paras.append((text, WHITE, True, 'en'))
    if cn: paras.append((cn, WHITE, False, 'cn'))
    return dict(fill=fill, paras=paras)

CN_ACT = {
    'Maintenance': '维护', 'WOT Implementation': 'WOT实施', 'Flight Test': '飞行测试',
    'Engine Run Test': '发动机试车', 'Select Activity': '选择活动',
}

def date_row(cw, dates, fri_black):
    cells = [L(f'CW {cw}')]
    for i, d in enumerate(dates):
        if fri_black and i == 4:
            cells.append(dict(fill=BLACK, paras=[(d, WHITE, True, 'en')]))
        else:
            cells.append(dict(fill=GREY, paras=[(d, BLACK, True, 'en')]))
    return dict(h=200000, cells=cells)

# ---------------- SN1003 (slide 7) ----------------
SN1003_DATES = ['Mo - 31.Aug.2026','Di - 01.Sep.2026','Mi - 02.Sep.2026','Do - 03.Sep.2026','Fr - 04.Sep.2026']
SN1003_AM = det([
    ('25FH Inspection (75%)\n100FH Inspection (75%)\n200FH Inspection (75%)', '25/100/200小时定检（75%）', RED),
    ('Main and Emergency Battery Maint (95%) (will be finished before flight)', '主电池及应急电池维护（95%）（飞行前完成）', RED),
    ('Scheduled Maintenance 1 Year Interval (50%)', '年度计划维护（50%）', RED),
])
SN1003_PM = det([
    ('FDR Accel Parameter - WOT 532 (80% - missing ADAU)', 'FDR加速度参数 - WOT 532（80%，缺少ADAU）', RED),
    ('TCAS Wiring Update - WOT 536 (80%- missing ADAU)', 'TCAS接线更新 - WOT 536（80%，缺少ADAU）', RED),
    ('Air/water tightness - WOT 524 (0%)', '气密/水密性检查 - WOT 524（0%）', BLUE),
    ('EEM and MHAM Replacement - WOT 551 (waiting for material)', 'EEM和MHAM替换 - WOT 551（等待物料）', RED),
    ('FTE Seat Re-installation (0%, will be finished after inspection)', 'FTE座椅重新安装（0%，将在检查后完成）', BLUE),
    ('Flap Actuator Harness Interface Change - WOT 556', '襟翼作动筒线束接口更改 - WOT 556', BLUE),
])
SN1003_RMK = det([
    ('Inspections are pending for ADAU installation and parts to be completed',
     '各项检查待ADAU安装及零件到位后进行', RED),
])
def sn1003_rows():
    return [
        dict(h=200000, cells=[L('CW 36')] + [dict(fill=GREY, paras=[(d, BLACK, True, 'en')]) for d in SN1003_DATES[:4]]
             + [dict(fill=BLACK, paras=[(SN1003_DATES[4], WHITE, True, 'en')])]),
        dict(h=190000, cells=[L('')] + [A_hdr('Maintenance', PURPLE)]*5),
        dict(h=800000, cells=[L('')] + [SN1003_AM]*5),
        dict(h=190000, cells=[L('')] + [A_hdr('WOT Implementation', PURPLE)]*5),
        dict(h=1500000, cells=[L('')] + [SN1003_PM]*5),
        dict(h=320000, cells=[L('RMK')] + [SN1003_RMK]*5),
    ]

# ---------------- SN1004 CW36 (slide 8) ----------------
SN1004_36_AM = [
    det([('4th Seat Installation', '第4座椅安装', BLUE),
         ('Engine Cooling Bypass Removal', '发动机冷却旁通堵头拆除', BLUE),
         ('Load Bank Removal', '负载箱拆除', BLUE)]),
    det([('Reserved for EASA Flight Readiness', '预留用于EASA飞行准备', BLUE)]),
    det([('EASA Panel 1 Pilot Familiarization Flight\n[MAS]\nLFTE:TRK', 'EASA Panel 1飞行员熟悉飞行', BLACK)]),
    det([('4th Seat Removal', '第4座椅拆除', BLUE),
         ('Ballast Implementation (TBD)', '配重实施（待定）', RED),
         ('Aircraft Release', '飞机放行', BLUE)]),
    det([('Build up Crosswind Campaign\n[MAS/KON]\nLFTE:STP', '侧风科目建立', BLACK)]),
]
SN1004_36_PM = [
    det([('Engine Run for Engine Cooling Bypass Removal', '发动机冷却旁通堵头拆除试车', BLACK)]),
    det([('EASA\nPilot Briefing and\nA/C Training', ['飞行员简报及', 'A/C培训'], BLUE)]),
    det([('EASA Panel 1 Pilot Familiarization Flight (BACKUP)\n[MAS]\nLFTE:TRK', 'EASA Panel 1飞行员熟悉飞行（备份）', BLACK)]),
    det([('Build up Crosswind Campaign\n[MAS/KON]\nLFTE:STP', '侧风科目建立', RED)]),
    det([('Build up Crosswind Campaign\n[MAS/KON]\nLFTE:STP', '侧风科目建立', BLACK)]),
]
def sn1004_36_rows():
    am_act = [A_hdr('WOT Implementation', PURPLE), A_hdr('Select Activity', SELGREY),
              A_hdr('Flight Test', ORANGE), A_hdr('WOT Implementation', PURPLE), A_hdr('Flight Test', ORANGE)]
    pm_act = [A_hdr('Engine Run Test', ORANGE), A_hdr('Select Activity', SELGREY),
              A_hdr('Flight Test', ORANGE), A_hdr('Flight Test', ORANGE), A_hdr('Flight Test', ORANGE)]
    blank = dict(fill=WHITE, paras=[])
    am_bar = [blank, blank, bar('Flight Test Executed', '飞行测试已执行', GREEN), blank,
              bar('Flight Test Executed', '飞行测试已执行', GREEN)]
    pm_bar = [bar('Engine Run Executed', '发动机试车已执行', GREEN), blank,
              bar('Flight Test Executed', '飞行测试已执行', GREEN), bar('', None, RED),
              bar('Flight Test Executed', '飞行测试已执行', GREEN)]
    rmk = [blank, blank, blank,
           det([('Flight is cancelled for EASA Activity', '因EASA活动取消飞行', RED)]), blank]
    return [
        date_row(36, SN1003_DATES, fri_black=True),
        dict(h=190000, cells=[L('AM')] + am_act),
        dict(h=620000, cells=[L('')] + SN1004_36_AM),
        dict(h=260000, cells=[L('')] + am_bar),
        dict(h=190000, cells=[L('PM')] + pm_act),
        dict(h=500000, cells=[L('')] + SN1004_36_PM),
        dict(h=260000, cells=[L('')] + pm_bar),
        dict(h=260000, cells=[L('RMK')] + rmk),
    ]

# ---------------- SN1004 CW37 (new slide 9) ----------------
SN1004_37_DATES = ['Mo - 07.Sep.2026','Di - 08.Sep.2026','Mi - 09.Sep.2026','Do - 10.Sep.2026','Fr - 11.Sep.2026']
SN1004_37_AM = [
    det([('[09:00] Engine Starvation Test', '[09:00] 发动机断油测试', BLUE),
         ('[11:00] Crosswind Campaign', '[11:00] 侧风科目', BLUE),
         ('Test Readiness Review', '测试就绪评审', BLUE)]),
    det([('Build up Crosswind Campaign\n[MAS/THE]\nLFTE:STP (not on board)', '侧风科目建立', BLUE)]),
    det([('Stall Speed Determination\n[MAS/THE]\nLFTE: HTR', '失速速度确定', BLUE)]),
    det([('Stall Speed Determination\n[MAS/THE]\nLFTE: MCS', '失速速度确定', BLUE)]),
    det([('FTE Rack Removal', 'FTE设备架拆除', BLUE)]),
]
SN1004_37_PM = [
    det([('Build up Crosswind Campaign\n[MAS/THE]\nLFTE:HTR', '侧风科目建立', BLUE)]),
    det([('Build up Crosswind Campaign\n[MAS/THE]\nLFTE:STP (not on board)', '侧风科目建立', BLUE)]),
    det([('Stall Speed Determination\n[MAS/THE]\nLFTE: TRK', '失速速度确定', BLUE)]),
    det([('FTE Rack Removal', 'FTE设备架拆除', BLUE)]),
    det([('Weighing', '称重', BLUE), ('A/C Readiness', '飞机准备状态', BLUE)]),
]
def sn1004_37_rows():
    am_act = [A_hdr('Engine Run Test', ORANGE), A_hdr('Flight Test', ORANGE), A_hdr('Flight Test', ORANGE),
              A_hdr('Flight Test', ORANGE), A_hdr('WOT Implementation', PURPLE)]
    pm_act = [A_hdr('Flight Test', ORANGE), A_hdr('Flight Test', ORANGE), A_hdr('Flight Test', ORANGE),
              A_hdr('WOT Implementation', PURPLE), A_hdr('Maintenance', PURPLE)]
    rmk1 = det([('After flight balasts will be removed', '飞行后拆除配重', BLACK)])
    blank = dict(fill=WHITE, paras=[])
    return [
        date_row(37, SN1004_37_DATES, fri_black=False),
        dict(h=190000, cells=[L('AM')] + am_act),
        dict(h=620000, cells=[L('')] + SN1004_37_AM),
        dict(h=190000, cells=[L('PM')] + pm_act),
        dict(h=500000, cells=[L('')] + SN1004_37_PM),
        dict(h=260000, cells=[L('RMK')] + [rmk1, blank, blank, blank, blank]),
    ]

def add_cn_note(slide, shape_name, cn_text):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name == shape_name:
            tb = sh.text_frame._txBody
            p = etree.Element(A+'p')
            pPr = etree.SubElement(p, A+'pPr'); pPr.set('algn','ctr')
            r = etree.SubElement(p, A+'r')
            rPr = etree.SubElement(r, A+'rPr')
            rPr.set('lang','zh-CN'); rPr.set('altLang','en-US'); rPr.set('sz','1000'); rPr.set('dirty','0')
            sf = etree.SubElement(rPr, A+'solidFill'); etree.SubElement(sf, A+'srgbClr').set('val','5F6F82')
            ea = etree.SubElement(rPr, A+'ea'); ea.set('typeface','Arial'); ea.set('pitchFamily','34'); ea.set('charset','0')
            etree.SubElement(r, A+'t').text = cn_text
            tb.append(p)
            return True
    return False

def set_title(slide, text):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.name.startswith('Title'):
            sh.text_frame.text = text
            return

def pass_a(inp, out):
    prs = Presentation(inp)
    s7, s8 = prs.slides[6], prs.slides[7]
    # remove table screenshots (keep banner strips: banners are wide but short)
    for sl in (s7, s8):
        for sh in list(sl.shapes):
            if sh.shape_type == 13 and sh.height > Emu(int(2*914400)):  # table imgs are 3.5-4.0in tall
                sh._element.getparent().remove(sh._element)
    build_table(s7, sn1003_rows(), int(0.37*914400), int(1.90*914400), int(12.50*914400), int(0.63*914400))
    build_table(s8, sn1004_36_rows(), int(0.37*914400), int(1.90*914400), int(12.44*914400), int(0.50*914400))
    set_title(s8, 'Short Term Flight Test Plan SN1004 (CW36)')
    add_cn_note(s7, 'Rectangle 7', '年检活动优先级较低，因ADAU维修导致强制停飞')
    prs.save(out)
    print('pass A done: slide7 SN1003 table, slide8 CW36 table')

def clone_slide8(src, out):
    """XML surgery: duplicate slide8 as a new last slide (slide9)."""
    work = os.path.join(os.path.dirname(os.path.abspath(out)), '_clone_work')
    if os.path.exists(work): shutil.rmtree(work)
    with zipfile.ZipFile(src) as z: z.extractall(work)
    shutil.copy(os.path.join(work,'ppt','slides','slide8.xml'),
                os.path.join(work,'ppt','slides','slide9.xml'))
    # rels: drop notesSlide reference
    rels_path = os.path.join(work,'ppt','slides','_rels','slide8.xml.rels')
    tree = etree.parse(rels_path)
    for rel in list(tree.getroot()):
        if rel.get('Type').endswith('/notesSlide'):
            tree.getroot().remove(rel)
    tree.write(os.path.join(work,'ppt','slides','_rels','slide9.xml.rels'),
               xml_declaration=True, encoding='UTF-8', standalone=True)
    # content types
    CT = 'http://schemas.openxmlformats.org/package/2006/content-types'
    ct_path = os.path.join(work,'[Content_Types].xml')
    ct = etree.parse(ct_path)
    ov = etree.SubElement(ct.getroot(), f'{{{CT}}}Override')
    ov.set('PartName','/ppt/slides/slide9.xml')
    ov.set('ContentType','application/vnd.openxmlformats-officedocument.presentationml.slide+xml')
    ct.write(ct_path, xml_declaration=True, encoding='UTF-8', standalone=True)
    # presentation rels
    pr_path = os.path.join(work,'ppt','_rels','presentation.xml.rels')
    pr = etree.parse(pr_path)
    rids = [int(re.sub(r'\D','',r.get('Id'))) for r in pr.getroot()]
    new_rid = f'rId{max(rids)+1}'
    rel = etree.SubElement(pr.getroot(), pr.getroot()[0].tag)
    rel.set('Id', new_rid)
    rel.set('Type', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide')
    rel.set('Target', 'slides/slide9.xml')
    pr.write(pr_path, xml_declaration=True, encoding='UTF-8', standalone=True)
    # sldIdLst
    pm_path = os.path.join(work,'ppt','presentation.xml')
    pm = etree.parse(pm_path)
    NS = {'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
          'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    lst = pm.getroot().find('p:sldIdLst', NS)
    ids = [int(e.get('id')) for e in lst]
    sld = etree.SubElement(lst, '{http://schemas.openxmlformats.org/presentationml/2006/main}sldId')
    sld.set('id', str(max(ids)+1))
    sld.set('{%s}id' % NS['r'], new_rid)
    pm.write(pm_path, xml_declaration=True, encoding='UTF-8', standalone=True)
    # repack
    if os.path.exists(out): os.remove(out)
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as zf:
        for root, dirs, files in os.walk(work):
            for f in files:
                fp = os.path.join(root,f)
                zf.write(fp, os.path.relpath(fp, work))
    shutil.rmtree(work)
    print(f'cloned slide8 -> slide9 ({new_rid})')

def pass_b(inp, out):
    prs = Presentation(inp)
    s9 = prs.slides[8]
    assert len(prs.slides) == 9
    # drop the cloned CW36 table, build CW37 table
    for sh in list(s9.shapes):
        if sh.has_table:
            sh._element.getparent().remove(sh._element)
    build_table(s9, sn1004_37_rows(), int(0.37*914400), int(1.90*914400), int(12.44*914400), int(0.50*914400))
    set_title(s9, 'Short Term Flight Test Plan SN1004 (CW37)')
    prs.save(out)
    print('pass B done: slide9 CW37 table')

if __name__ == '__main__':
    inp, out = sys.argv[1], sys.argv[2]
    tmp1 = out + '.step1.pptx'
    pass_a(inp, tmp1)
    clone_slide8(tmp1, out + '.step2.pptx')
    pass_b(out + '.step2.pptx', out)
    os.remove(tmp1); os.remove(out + '.step2.pptx')
    print(f'Done: {out}')
